import io
import json
import re
import sys
from pathlib import Path

import openpyxl
from PIL import Image


CATEGORY_NAMES = {
    "rank": "Rank",
    "uniform": "Uniform",
    "primary": "Primary",
    "secondary": "Secondary",
    "medical": "Medical",
    "grenades": "Grenades",
    "ammo": "Ammo",
    "equipment": "Equipment",
}

VERTICAL_SHEETS = {
    "Recruits": {"label": "Recruits", "description": "Recruit and cadet uniforms, equipment, and approved loadouts."},
    "Infantry": {"label": "Infantry", "description": "Standard infantry uniforms and loadouts by rank."},
    "Certifications": {"label": "Certifications", "description": "Certification-specific uniforms, weapons, equipment, and loadout limits."},
    "Hammer": {"label": "Hammer", "description": "Hammer aviation uniforms and loadouts by flight rank."},
    "Longbow": {"label": "Longbow", "description": "Longbow specialist uniforms and operational loadouts."},
}

SECTION_ORDER = [
    ("restricted-items", "Restricted Items", "Items prohibited or restricted during Main Operations."),
    ("recruits", "Recruits", VERTICAL_SHEETS["Recruits"]["description"]),
    ("infantry", "Infantry", VERTICAL_SHEETS["Infantry"]["description"]),
    ("certifications", "Certifications", VERTICAL_SHEETS["Certifications"]["description"]),
    ("dagger", "Dagger", "Dagger airborne uniforms and equipment by billet and rank."),
    ("hammer", "Hammer", VERTICAL_SHEETS["Hammer"]["description"]),
    ("zeus", "Zeus", "Approved Zeus appearance standards by qualification level."),
    ("arc", "ARC", "ARC uniform reference and supporting documentation."),
    ("longbow", "Longbow", VERTICAL_SHEETS["Longbow"]["description"]),
]


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def slug(value):
    result = re.sub(r"[^a-z0-9]+", "-", clean(value).lower()).strip("-")
    return result[:80] or "entry"


def cell_item(cell):
    text = clean(cell.value)
    if not text:
        return None
    url = cell.hyperlink.target if cell.hyperlink else None
    return {"text": text, **({"url": url} if url else {})}


def unique_key(section, title, used):
    base = f"{section}-{slug(title)}"
    key = base
    index = 2
    while key in used:
        key = f"{base}-{index}"
        index += 1
    used.add(key)
    return key


def export_image(image, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(io.BytesIO(image._data())) as source:
        source.thumbnail((1400, 1400), Image.Resampling.LANCZOS)
        if source.mode not in ("RGB", "RGBA"):
            source = source.convert("RGBA" if "transparency" in source.info else "RGB")
        source.save(destination, "WEBP", quality=84, method=6)


def find_header(sheet):
    for row in sheet.iter_rows():
        normalized = [clean(cell.value).lower() for cell in row]
        if "rank" in normalized and sum(name in normalized for name in CATEGORY_NAMES) >= 5:
            columns = {}
            for index, value in enumerate(normalized, 1):
                if value in CATEGORY_NAMES:
                    columns[CATEGORY_NAMES[value]] = index
            return row[0].row, columns
    raise ValueError(f"No category header found for {sheet.title}")


def image_for_range(sheet, start, end):
    for image in sheet._images:
        row = image.anchor._from.row + 1
        if start <= row < end:
            return image
    return None


def vertical_entries(sheet, section, assets_root, used):
    header_row, columns = find_header(sheet)
    title_col = columns.pop("Rank")
    starts = []
    for row in range(header_row + 1, sheet.max_row + 1):
        title = clean(sheet.cell(row, title_col).value)
        populated = sum(bool(clean(sheet.cell(row, column).value)) for column in columns.values())
        if title and populated >= 2:
            starts.append(row)

    entries = []
    for order, start in enumerate(starts):
        end = starts[order + 1] if order + 1 < len(starts) else sheet.max_row + 1
        title = clean(sheet.cell(start, title_col).value)
        categories = []
        for label, column in columns.items():
            items = []
            for row in range(start, end):
                item = cell_item(sheet.cell(row, column))
                if item and item not in items:
                    items.append(item)
            if items:
                categories.append({"label": label, "items": items})

        notes = []
        links = []
        for row in range(start + 1, end):
            item = cell_item(sheet.cell(row, title_col))
            if not item:
                continue
            if item.get("url"):
                links.append(item)
            elif "kit import" in item["text"].lower() or item["text"].startswith("|") or item["text"].startswith("*") or item["text"].lower().startswith("note"):
                notes.append(item["text"])

        key = unique_key(section, title, used)
        image_path = None
        image = image_for_range(sheet, start, end)
        if image:
            relative = Path("uniform-guide") / section / f"{key}.webp"
            export_image(image, assets_root / relative)
            image_path = f"/{relative.as_posix()}"

        entries.append({
            "entry_key": key,
            "section_key": section,
            "title": title,
            "subtitle": "",
            "summary": "",
            "categories": categories,
            "notes": notes,
            "links": links,
            "image_path": image_path,
            "sort_order": order,
        })
    return entries


def restricted_entries(sheet, assets_root, used):
    header_row = 4
    header_columns = [(cell.column, clean(cell.value)) for cell in sheet[header_row] if clean(cell.value)]
    categories = []
    for column, label in header_columns:
        items = []
        for row in range(header_row + 1, sheet.max_row + 1):
            item = cell_item(sheet.cell(row, column))
            if item:
                items.append(item)
        if items:
            categories.append({"label": label, "items": items})
    return [{
        "entry_key": unique_key("restricted-items", "Restricted Main Operation Items", used),
        "section_key": "restricted-items",
        "title": "Restricted Main Operation Items",
        "subtitle": "Main Operations",
        "summary": "Unless explicitly approved otherwise, the listed items must not be used during Main Operations.",
        "categories": categories,
        "notes": [clean(sheet["B3"].value)] if clean(sheet["B3"].value) else [],
        "links": [],
        "image_path": None,
        "sort_order": 0,
    }]


def matrix_entries(sheet, section, header_row, first_data_row, columns, category_labels, assets_root, used):
    entries = []
    images_by_column = {}
    for image in sheet._images:
        images_by_column[image.anchor._from.col + 1] = image
    category_rows = {}
    for row in range(first_data_row, sheet.max_row + 1):
        values = {clean(sheet.cell(row, column).value).lower().rstrip(":") for column in range(1, sheet.max_column + 1)}
        matched = next((label for label in category_labels if label.lower() in values), None)
        if matched:
            category_rows[row] = "Ammunition & Grenades" if matched == "Ammuntion & Grenades" else matched

    for order, column in enumerate(columns):
        title = clean(sheet.cell(header_row, column).value)
        if not title:
            continue
        categories = []
        current_label = "Standard"
        current_items = []
        for row in range(first_data_row, sheet.max_row + 1):
            matched = category_rows.get(row)
            if matched:
                if current_items:
                    categories.append({"label": current_label, "items": current_items})
                current_label = matched
                current_items = []
                continue
            value = clean(sheet.cell(row, column).value)
            if not value:
                continue
            item = cell_item(sheet.cell(row, column))
            if item:
                current_items.append(item)
        if current_items:
            categories.append({"label": current_label, "items": current_items})

        key = unique_key(section, title, used)
        image_path = None
        image = images_by_column.get(column)
        if image:
            relative = Path("uniform-guide") / section / f"{key}.webp"
            export_image(image, assets_root / relative)
            image_path = f"/{relative.as_posix()}"
        entries.append({
            "entry_key": key,
            "section_key": section,
            "title": title,
            "subtitle": "",
            "summary": "",
            "categories": categories,
            "notes": [],
            "links": [],
            "image_path": image_path,
            "sort_order": order,
        })
    return entries


def arc_entries(sheet, used):
    cell = sheet["A2"]
    link = cell.hyperlink.target if cell.hyperlink else None
    return [{
        "entry_key": unique_key("arc", "ARC Uniform Standards", used),
        "section_key": "arc",
        "title": "ARC Uniform Standards",
        "subtitle": "External reference",
        "summary": "The source workbook currently contains a broken ARC reference. Use the linked ARC document until this section is replaced with structured standards.",
        "categories": [],
        "notes": ["Source workbook cell A1 currently returns #REF!"],
        "links": ([{"text": clean(cell.value) or "ARC Uniform Document", "url": link}] if link else []),
        "image_path": None,
        "sort_order": 0,
    }]


def sql_literal(value):
    return "'" + value.replace("'", "''") + "'"


def build_sql(sections, entries):
    lines = []
    for order, section in enumerate(sections):
        lines.append(
            "(" + ", ".join([
                sql_literal(section[0]), sql_literal(section[1]), sql_literal(section[2]), str(order)
            ]) + ")"
        )
    section_values = ",\n  ".join(lines)

    entry_lines = []
    for entry in entries:
        values = [
            sql_literal(entry["entry_key"]),
            sql_literal(entry["section_key"]),
            sql_literal(entry["title"]),
            sql_literal(entry["subtitle"]),
            sql_literal(entry["summary"]),
            sql_literal(json.dumps(entry["categories"], ensure_ascii=False, separators=(",", ":"))) + "::jsonb",
            "array[" + ",".join(sql_literal(note) for note in entry["notes"]) + "]::text[]" if entry["notes"] else "'{}'::text[]",
            sql_literal(json.dumps(entry["links"], ensure_ascii=False, separators=(",", ":"))) + "::jsonb",
            sql_literal(entry["image_path"]) if entry["image_path"] else "null",
            str(entry["sort_order"]),
        ]
        entry_lines.append("(" + ", ".join(values) + ")")

    return section_values, ",\n  ".join(entry_lines)


def main():
    if len(sys.argv) != 4:
        raise SystemExit("Usage: generate_seed.py <source.xlsx> <public-dir> <output.json>")
    workbook_path = Path(sys.argv[1])
    public_dir = Path(sys.argv[2])
    output_path = Path(sys.argv[3])
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    sheets = {sheet.title.strip(): sheet for sheet in workbook.worksheets}
    used = set()
    entries = []
    entries.extend(restricted_entries(sheets["Restricted Items"], public_dir, used))
    for name, config in VERTICAL_SHEETS.items():
        entries.extend(vertical_entries(sheets[name], slug(name), public_dir, used))
    entries.extend(matrix_entries(
        sheets["Dagger"], "dagger", 2, 24, [2, 3, 4, 6, 7, 9, 10],
        ["Arsenal Kit Import", "Trooper Loadout", "Helmet", "Uniform", "Jumppack", "Vest", "Medical (See Medical Loadouts)", "Equipment", "Ammuntion & Grenades"],
        public_dir, used,
    ))
    entries.extend(matrix_entries(
        sheets["Zeus"], "zeus", 5, 6, [2, 3, 4, 5],
        ["Headwear", "Uniform", "Jumppack"], public_dir, used,
    ))
    entries.extend(arc_entries(sheets["ARC"], used))
    entries.sort(key=lambda entry: ([section[0] for section in SECTION_ORDER].index(entry["section_key"]), entry["sort_order"]))
    section_values, entry_values = build_sql(SECTION_ORDER, entries)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps({
        "sections": [
            {"key": key, "label": label, "description": description, "sortOrder": order}
            for order, (key, label, description) in enumerate(SECTION_ORDER)
        ],
        "entries": entries,
        "sectionSqlValues": section_values,
        "entrySqlValues": entry_values,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"sections": len(SECTION_ORDER), "entries": len(entries), "images": len(list((public_dir / "uniform-guide").rglob("*.webp"))) }))


if __name__ == "__main__":
    main()
